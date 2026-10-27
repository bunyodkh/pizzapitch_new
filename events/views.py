import csv
import json
import threading
from datetime import datetime, timedelta

from django.conf import settings
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.views import LoginView, LogoutView
from django.core.exceptions import ValidationError
from django.core.serializers.json import DjangoJSONEncoder
from django.core.validators import validate_email
from django.db.models import Count
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils.dateformat import format as format_date
from django.utils.safestring import mark_safe
from django.utils.translation import gettext as _
from django.views.decorators.http import require_GET, require_http_methods, require_POST

from .forms import EventForm, ExpertForm, PartnerForm, TeamAuthenticationForm
from .models import Application, Event, Expert, Partner, PartnerInquiry, Subscriber
from .telegram_notify import (
    notify_new_application_telegram,
    notify_new_partner_inquiry_telegram,
    notify_new_subscriber_telegram,
)

# Pseudo applicant-type value for the contacts page's "type" filter — picks
# out waitlist Subscribers, which aren't Applications and have no real type.
CONTACT_TYPE_SUBSCRIBER = 'subscriber'


def current_event():
    return Event.objects.filter(is_active=True).order_by('date', 'time').first()


def event_jsonld(event):
    """Schema.org Event structured data for the landing page — lets search
    engines show the event as a rich result (Google Events, "things to do").
    """
    location = {'@type': 'Place', 'name': event.place}
    if event.city:
        location['address'] = {'@type': 'PostalAddress', 'addressLocality': event.city, 'addressCountry': 'UZ'}
    if event.location_url:
        location['hasMap'] = event.location_url

    data = {
        '@context': 'https://schema.org',
        '@type': 'Event',
        'name': event.title,
        'startDate': datetime.combine(event.date, event.time).isoformat(),
        'eventAttendanceMode': 'https://schema.org/OfflineEventAttendanceMode',
        'eventStatus': 'https://schema.org/EventScheduled',
        'description': event.description or event.tagline or event.title,
        'location': location,
        'organizer': {'@type': 'Organization', 'name': 'PizzaPitch', 'url': settings.SITE_URL},
        'offers': {
            '@type': 'Offer',
            'price': str(event.guest_price),
            'priceCurrency': event.currency,
            'availability': (
                'https://schema.org/InStock' if event.guest_registration_open else 'https://schema.org/SoldOut'
            ),
            'url': settings.SITE_URL,
        },
    }
    if event.hero_photo:
        data['image'] = [settings.SITE_URL + event.hero_photo.url]

    # Escape </script>-breaking sequences the same way Django's json_script does —
    # description/title are staff-entered, not user input, but this is free insurance.
    raw = json.dumps(data, cls=DjangoJSONEncoder)
    raw = raw.replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026')
    return mark_safe(raw)


@require_GET
def index(request):
    event = current_event()
    if event is None:
        notify_i18n = {
            'invalidEmail': _('Enter a valid email'),
            'success': _("Done! We'll email you as soon as registration opens."),
            'error': _('Could not submit. Please try again.'),
        }
        return render(request, 'between.html', {'notify_i18n': notify_i18n})

    price = f'{event.startup_price:.0f}'
    event_date = format_date(event.date, 'D, d M')

    # "How the evening runs" times, relative to the event's own start time,
    # so they stay correct if the start time ever changes instead of drifting
    # like before. Doors at start; pitches 30 min in; networking starts once
    # the pitch block ends — 10 startups × (3 min pitch + 7 min Q&A) = 100 min.
    PITCH_BLOCK_MINUTES = 10 * (3 + 7)
    start_dt = datetime.combine(event.date, event.time)
    pitches_dt = start_dt + timedelta(minutes=30)
    schedule = {
        'doors': event.time.strftime('%H:%M'),
        'pitches': pitches_dt.strftime('%H:%M'),
        'networking': (pitches_dt + timedelta(minutes=PITCH_BLOCK_MINUTES)).strftime('%H:%M'),
    }

    i18n_strings = {
        'formTitleStartup': _('Startup application'),
        'formTitleGuest': _('Guest application'),
        'formSubStartup': _('A short note about your project. Selected teams get the pitch format.'),
        'formSubGuest': _('A couple fields, we will follow up with confirmation.'),
        'doneSubTemplate': _('Thanks! Confirmation will be sent to {email}.'),
        'doneStepsStartup': [
            _('We review applications manually, usually 24–48 hours.'),
            _('Approved: we email a payment link for participation ({price} {currency}).').format(
                price=price, currency=event.currency,
            ),
            _('Paid: your pitch slot is locked in. Prepare your 3-minute pitch.'),
        ],
        'doneStepsGuest': [
            _('We review applications manually, usually 24–48 hours.'),
            _('We email you a confirmation with all the evening details.'),
            _('Confirmed, add it to your calendar. See you on {date}.').format(date=event_date),
        ],
        'fillFieldsPrefix': _('Please fill in: '),
        'genericError': _('Could not submit the application. Check your connection and try again.'),
    }
    if event.landing_page == Event.LANDING_WOMEN:
        # PizzaPitch Women is free for everyone, no payment step after approval.
        i18n_strings['doneStepsStartup'] = [
            _('We review applications manually, usually 24–48 hours.'),
            _('Selected: your team is in the semifinal. We email you all the details.'),
            _('Prepare your pitch: 5 minutes + 5 minutes of Q&A.'),
        ]
        i18n_strings['doneStepsGuest'][2] = _('Confirmed, add it to your calendar. See you on {date}.').format(
            date=_('24 Nov 2026'),
        )

    experts = Expert.objects.filter(is_visible=True).order_by('order', 'id')
    partners = Partner.objects.filter(is_visible=True).order_by('order', 'id')
    template = 'index_women.html' if event.landing_page == Event.LANDING_WOMEN else 'index.html'
    return render(request, template, {
        'event': event, 'schedule': schedule, 'experts': experts, 'partners': partners, 'i18n_strings': i18n_strings,
        'event_jsonld': event_jsonld(event),
    })


@require_POST
def apply(request):
    event = current_event()
    if event is None:
        return JsonResponse({'error': _('There is no open event to apply to.')}, status=404)

    applicant_type = request.POST.get('type', '')
    if applicant_type not in (Application.STARTUP, Application.GUEST):
        return JsonResponse({'error': _('Invalid application type.')}, status=400)

    if applicant_type == Application.STARTUP and not event.startup_registration_open:
        return JsonResponse({'error': _('Startup registration is currently closed.')}, status=403)
    if applicant_type == Application.GUEST and not event.guest_registration_open:
        return JsonResponse({'error': _('Guest registration is currently closed.')}, status=403)

    name = request.POST.get('name', '').strip()
    email = request.POST.get('email', '').strip()
    telegram = request.POST.get('tg', '').strip()
    if not name or not email or not telegram:
        return JsonResponse({'error': _('Name, email and Telegram are required.')}, status=400)

    if applicant_type == Application.STARTUP:
        project = request.POST.get('project', '').strip()
        pitch = request.POST.get('pitch', '').strip()
        phone = request.POST.get('phone', '').strip()
        if not project or not pitch or not phone:
            return JsonResponse({'error': _('Project, phone and pitch are required.')}, status=400)
    else:
        company = request.POST.get('company', '').strip()
        why = request.POST.get('why', '').strip()
        if not company or not why:
            return JsonResponse({'error': _('Company and reason are required.')}, status=400)

        guest_count = Application.objects.filter(
            event=event, applicant_type=Application.GUEST,
        ).exclude(status=Application.STATUS_REJECTED).count()
        if guest_count >= Application.GUEST_CAP:
            return JsonResponse({'error': _('Guest spots are full, no seats left.')}, status=409)

    application = Application.objects.create(
        event=event,
        applicant_type=applicant_type,
        name=name,
        email=email,
        phone=request.POST.get('phone', '').strip(),
        telegram=telegram,
        project=request.POST.get('project', '').strip(),
        link=request.POST.get('link', '').strip(),
        pitch=request.POST.get('pitch', '').strip(),
        company=request.POST.get('company', '').strip(),
        why=request.POST.get('why', '').strip(),
    )
    # Fire-and-forget: Telegram delivery has its own retry/backoff and can take
    # several seconds on network trouble, so it must not hold up this response.
    threading.Thread(target=notify_new_application_telegram, args=(application,), daemon=True).start()
    return JsonResponse({'id': application.id, 'status': application.status}, status=201)


@require_POST
def subscribe(request):
    email = request.POST.get('email', '').strip()
    if not email:
        return JsonResponse({'error': _('Email is required.')}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return JsonResponse({'error': _('Invalid email.')}, status=400)

    subscriber, created = Subscriber.objects.get_or_create(email=email)
    if created:
        threading.Thread(target=notify_new_subscriber_telegram, args=(subscriber,), daemon=True).start()
    return JsonResponse({'status': 'subscribed'}, status=201 if created else 200)


@require_POST
def partner_inquiry(request):
    email = request.POST.get('email', '').strip()
    message = request.POST.get('message', '').strip()
    if not email or not message:
        return JsonResponse({'error': _('Email and message are required.')}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return JsonResponse({'error': _('Invalid email.')}, status=400)

    inquiry = PartnerInquiry.objects.create(email=email, message=message)
    threading.Thread(target=notify_new_partner_inquiry_telegram, args=(inquiry,), daemon=True).start()
    return JsonResponse({'status': 'received'}, status=201)


class TeamLoginView(LoginView):
    template_name = 'team_login.html'
    authentication_form = TeamAuthenticationForm
    redirect_authenticated_user = True


class TeamLogoutView(LogoutView):
    next_page = 'team_login'


@staff_member_required(login_url='team_login')
@require_GET
def team_home(request):
    return render(request, 'team_home.html', {'active_tab': 'home'})


@staff_member_required(login_url='team_login')
@require_GET
def event_list(request):
    events = Event.objects.order_by('-date', '-time')
    return render(request, 'event_list.html', {'events': events, 'active_tab': 'events'})


@staff_member_required(login_url='team_login')
@require_http_methods(['GET', 'POST'])
def event_create(request):
    if request.method == 'POST':
        form = EventForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect('event_list')
    else:
        form = EventForm()
    return render(request, 'event_form.html', {'form': form, 'event': None, 'active_tab': 'events'})


@staff_member_required(login_url='team_login')
@require_http_methods(['GET', 'POST'])
def event_edit(request, event_id):
    event = get_object_or_404(Event, pk=event_id)
    if request.method == 'POST':
        form = EventForm(request.POST, request.FILES, instance=event)
        if form.is_valid():
            obj = form.save(commit=False)
            if form.cleaned_data.get('remove_hero_photo') and not request.FILES.get('hero_photo') and obj.hero_photo:
                obj.hero_photo.delete(save=False)
            obj.save()
            return redirect('event_list')
    else:
        form = EventForm(instance=event)
    return render(request, 'event_form.html', {'form': form, 'event': event, 'active_tab': 'events'})


@staff_member_required(login_url='team_login')
@require_POST
def event_delete(request, event_id):
    event = get_object_or_404(Event, pk=event_id)
    event.delete()
    return redirect('event_list')


@staff_member_required(login_url='team_login')
@require_GET
def expert_list(request):
    experts = Expert.objects.order_by('order', 'id')
    return render(request, 'expert_list.html', {'experts': experts, 'active_tab': 'experts'})


@staff_member_required(login_url='team_login')
@require_http_methods(['GET', 'POST'])
def expert_create(request):
    if request.method == 'POST':
        form = ExpertForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect('expert_list')
    else:
        form = ExpertForm()
    return render(request, 'expert_form.html', {'form': form, 'expert': None, 'active_tab': 'experts'})


@staff_member_required(login_url='team_login')
@require_http_methods(['GET', 'POST'])
def expert_edit(request, expert_id):
    expert = get_object_or_404(Expert, pk=expert_id)
    if request.method == 'POST':
        form = ExpertForm(request.POST, request.FILES, instance=expert)
        if form.is_valid():
            form.save()
            return redirect('expert_list')
    else:
        form = ExpertForm(instance=expert)
    return render(request, 'expert_form.html', {'form': form, 'expert': expert, 'active_tab': 'experts'})


@staff_member_required(login_url='team_login')
@require_POST
def expert_delete(request, expert_id):
    expert = get_object_or_404(Expert, pk=expert_id)
    expert.delete()
    return redirect('expert_list')


@staff_member_required(login_url='team_login')
@require_GET
def partner_list(request):
    partners = Partner.objects.order_by('order', 'id')
    return render(request, 'partner_list.html', {'partners': partners, 'active_tab': 'partners'})


@staff_member_required(login_url='team_login')
@require_http_methods(['GET', 'POST'])
def partner_create(request):
    if request.method == 'POST':
        form = PartnerForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            return redirect('partner_list')
    else:
        form = PartnerForm()
    return render(request, 'partner_form.html', {'form': form, 'partner': None, 'active_tab': 'partners'})


@staff_member_required(login_url='team_login')
@require_http_methods(['GET', 'POST'])
def partner_edit(request, partner_id):
    partner = get_object_or_404(Partner, pk=partner_id)
    if request.method == 'POST':
        form = PartnerForm(request.POST, request.FILES, instance=partner)
        if form.is_valid():
            form.save()
            return redirect('partner_list')
    else:
        form = PartnerForm(instance=partner)
    return render(request, 'partner_form.html', {'form': form, 'partner': partner, 'active_tab': 'partners'})


@staff_member_required(login_url='team_login')
@require_POST
def partner_delete(request, partner_id):
    partner = get_object_or_404(Partner, pk=partner_id)
    partner.delete()
    return redirect('partner_list')


def _participant_applications(request, events):
    """Shared event/type/status filtering for the applications dashboard and its exports.

    Returns the selected event plus both the type-filtered-only queryset
    (used for the status tiles, which must stay stable regardless of which
    status is selected) and the fully filtered queryset (event + type + status).
    """
    active = current_event()

    event_param = request.GET.get('event', '')
    type_param = request.GET.get('type', 'all')
    status_param = request.GET.get('status', 'all')

    if event_param == 'all':
        selected_event = None
    elif event_param:
        selected_event = events.filter(pk=event_param).first()
    else:
        selected_event = active

    effective_event_param = event_param or (str(active.id) if active else '')

    base_applications = Application.objects.all()
    if selected_event is not None:
        base_applications = base_applications.filter(event=selected_event)
    elif event_param != 'all':
        # No active event and nothing explicitly selected — don't dump every event by accident.
        base_applications = base_applications.none()

    if type_param in (Application.STARTUP, Application.GUEST):
        base_applications = base_applications.filter(applicant_type=type_param)

    applications = base_applications.select_related('event')
    if status_param in (Application.STATUS_PENDING, Application.STATUS_APPROVED, Application.STATUS_REJECTED):
        applications = applications.filter(status=status_param)
    applications = applications.order_by('-created_at')

    return selected_event, effective_event_param, type_param, status_param, base_applications, applications


_APPLICATION_EXPORT_HEADERS = [
    'Name', 'Email', 'Phone', 'Telegram', 'Type', 'Status', 'Event', 'Submitted',
    'Project', 'Pitch deck', 'Pitch', 'Company', 'Why',
]


def _application_export_row(app):
    return [
        app.name, app.email, app.phone, app.telegram,
        app.get_applicant_type_display(), app.get_status_display(),
        app.event.title, app.created_at.strftime('%Y-%m-%d %H:%M'),
        app.project, app.link, app.pitch, app.company, app.why,
    ]


def _xlsx_response(filename, headers, rows):
    """Builds an .xlsx download from a header row and a list of row lists."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for i, header in enumerate(headers, start=1):
        column_values = [str(header)] + [str(row[i - 1]) for row in rows]
        width = max(len(v) for v in column_values)
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 50)
    for row in rows:
        ws.append(row)

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@staff_member_required(login_url='team_login')
@require_GET
def participants(request):
    """Participant management dashboard for the PizzaPitch team."""
    events = Event.objects.order_by('-date', '-time')
    (
        selected_event, effective_event_param, type_param, status_param,
        base_applications, applications,
    ) = _participant_applications(request, events)

    status_counts = {choice: 0 for choice, _ in Application.STATUS_CHOICES}
    for row in base_applications.values('status').annotate(n=Count('id')):
        status_counts[row['status']] = row['n']
    total_count = sum(status_counts.values())

    guest_count = None
    if selected_event is not None:
        guest_count = Application.objects.filter(
            event=selected_event, applicant_type=Application.GUEST,
        ).exclude(status=Application.STATUS_REJECTED).count()

    def qs_for_status(status_value):
        params = request.GET.copy()
        params['status'] = status_value
        params.setdefault('event', effective_event_param)
        params.setdefault('type', type_param)
        return params.urlencode()

    status_tiles = [
        {'key': 'all', 'label': _('All'), 'count': total_count, 'qs': qs_for_status('all')},
        {'key': Application.STATUS_PENDING, 'label': _('Pending'),
         'count': status_counts[Application.STATUS_PENDING], 'qs': qs_for_status(Application.STATUS_PENDING)},
        {'key': Application.STATUS_APPROVED, 'label': _('Approved'),
         'count': status_counts[Application.STATUS_APPROVED], 'qs': qs_for_status(Application.STATUS_APPROVED)},
        {'key': Application.STATUS_REJECTED, 'label': _('Rejected'),
         'count': status_counts[Application.STATUS_REJECTED], 'qs': qs_for_status(Application.STATUS_REJECTED)},
    ]

    context = {
        'events': events,
        'selected_event': selected_event,
        'event_param': effective_event_param,
        'type_param': type_param,
        'status_param': status_param,
        'applications': applications,
        'guest_count': guest_count,
        'guest_cap': Application.GUEST_CAP,
        'querystring': request.GET.urlencode(),
        'total_count': total_count,
        'status_tiles': status_tiles,
        'active_tab': 'applications',
    }
    return render(request, 'participants.html', context)


@staff_member_required(login_url='team_login')
@require_GET
def participants_export_csv(request):
    """CSV export of the currently filtered applications list."""
    events = Event.objects.order_by('-date', '-time')
    *_unused, applications = _participant_applications(request, events)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="applications.csv"'
    writer = csv.writer(response)
    writer.writerow(_APPLICATION_EXPORT_HEADERS)
    for app in applications:
        writer.writerow(_application_export_row(app))
    return response


@staff_member_required(login_url='team_login')
@require_GET
def participants_export_excel(request):
    """Excel export of the currently filtered applications list."""
    events = Event.objects.order_by('-date', '-time')
    *_unused, applications = _participant_applications(request, events)

    rows = [_application_export_row(app) for app in applications]
    return _xlsx_response('applications.xlsx', _APPLICATION_EXPORT_HEADERS, rows)


@staff_member_required(login_url='team_login')
@require_POST
def set_application_status(request, app_id):
    application = get_object_or_404(Application, pk=app_id)
    status = request.POST.get('status')
    if status in dict(Application.STATUS_CHOICES):
        application.status = status
        application.save(update_fields=['status'])

    url = reverse('participants')
    qs = request.POST.get('querystring', '')
    if qs:
        url = f'{url}?{qs}'
    return redirect(url)


def _filtered_applications(request, events):
    """Shared event/type/status filtering for the contacts list and its CSV export.

    Unlike the applications review queue, this is a contact directory, not
    an operational to-do list — with no event filter selected it should
    show every applicant on file, not just the current event's (there may
    be no current event at all between events, and past contacts still
    matter for outreach).
    """
    event_param = request.GET.get('event', '')
    type_param = request.GET.get('type', 'all')
    status_param = request.GET.get('status', 'all')

    selected_event = None
    if event_param and event_param != 'all':
        selected_event = events.filter(pk=event_param).first()
    effective_event_param = event_param or 'all'

    application_qs = Application.objects.all()
    if selected_event is not None:
        application_qs = application_qs.filter(event=selected_event)

    if type_param in (Application.STARTUP, Application.GUEST):
        application_qs = application_qs.filter(applicant_type=type_param)
    elif type_param == CONTACT_TYPE_SUBSCRIBER:
        # Subscribers aren't Applications at all — the "type" filter is
        # standing in for "which kind of contact", not an Application field.
        application_qs = application_qs.none()
    if status_param in (Application.STATUS_PENDING, Application.STATUS_APPROVED, Application.STATUS_REJECTED):
        application_qs = application_qs.filter(status=status_param)

    application_qs = application_qs.select_related('event')
    return selected_event, effective_event_param, type_param, status_param, application_qs


def _contact_rows(request):
    """Builds the combined contact directory: applicants plus waitlist
    subscribers. Subscribers aren't tied to an event/status, so they're only
    folded in when neither filter narrows the applicant list — otherwise
    "guests for event X" would misleadingly include people who never applied
    to anything. The type filter has a dedicated "Subscribers" option instead.
    """
    events = Event.objects.order_by('-date', '-time')
    selected_event, effective_event_param, type_param, status_param, application_qs = _filtered_applications(
        request, events,
    )

    rows = [
        {
            'name': app.name, 'email': app.email, 'phone': app.phone, 'telegram': app.telegram,
            'type_display': app.get_applicant_type_display(), 'status_display': app.get_status_display(),
            'event_title': app.event.title, 'source': _('Applicant'),
        }
        for app in application_qs
    ]

    include_subscribers = (
        selected_event is None and status_param == 'all'
        and type_param in ('all', CONTACT_TYPE_SUBSCRIBER)
    )
    if include_subscribers:
        rows += [
            {
                'name': '', 'email': sub.email, 'phone': '', 'telegram': '',
                'type_display': '', 'status_display': '', 'event_title': '', 'source': _('Subscriber'),
            }
            for sub in Subscriber.objects.all()
        ]

    rows.sort(key=lambda row: (row['name'] or row['email']).lower())
    return events, selected_event, effective_event_param, type_param, status_param, rows


@staff_member_required(login_url='team_login')
@require_GET
def contacts(request):
    """Contact directory — applicants and waitlist subscribers — for future outreach (e.g. event newsletters)."""
    events, selected_event, effective_event_param, type_param, status_param, rows = _contact_rows(request)

    context = {
        'events': events,
        'selected_event': selected_event,
        'event_param': effective_event_param,
        'type_param': type_param,
        'status_param': status_param,
        'contacts': rows,
        'total_count': len(rows),
        'querystring': request.GET.urlencode(),
        'active_tab': 'contacts',
    }
    return render(request, 'contacts.html', context)


_CONTACT_EXPORT_HEADERS = ['Name', 'Email', 'Phone', 'Telegram', 'Type', 'Status', 'Event', 'Source']


def _contact_export_row(row):
    return [
        row['name'], row['email'], row['phone'], row['telegram'],
        row['type_display'], row['status_display'], row['event_title'], row['source'],
    ]


@staff_member_required(login_url='team_login')
@require_GET
def contacts_export(request):
    """CSV export of the currently filtered contact list."""
    *_unused, rows = _contact_rows(request)

    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="contacts.csv"'
    writer = csv.writer(response)
    writer.writerow(_CONTACT_EXPORT_HEADERS)
    for row in rows:
        writer.writerow(_contact_export_row(row))
    return response


@staff_member_required(login_url='team_login')
@require_GET
def contacts_export_excel(request):
    """Excel export of the currently filtered contact list."""
    *_unused, rows = _contact_rows(request)
    return _xlsx_response('contacts.xlsx', _CONTACT_EXPORT_HEADERS, [_contact_export_row(row) for row in rows])
